#!/usr/bin/env python3
"""
make-filter-fixtures.py -- build the small .docx files that
run-filter-tests.py converts.

    python3 tests/make-filter-fixtures.py

The built files are committed, so running the tests needs nothing beyond
Pandoc. This script is for regenerating or extending them, and needs
python-docx (`pip3 install python-docx`).

WHAT THESE ARE FOR

Every bug in the v0.2 changelog under Fixed was silent: alt text applied
to the wrong image, two H1 elements on a page, a caption that stopped
being found, a table reduced to one empty cell. None of them raised an
error, and most were found by comparing two runs of a real book -- which
works only as long as the previous version is still installed. A fixture
carrying the structure that triggered the bug makes the check
self-contained.

Each file is deliberately small and carries several unrelated triggers,
because the cost is one Pandoc invocation per file rather than per
assertion.

WHAT WORD STORES THAT CATCHES PEOPLE OUT

Two of these are worth knowing if you extend the fixtures.

Pandoc's docx reader takes the document title and author from paragraphs
styled Title and Author, consuming them out of the body. Not from
docProps/core.xml, which it does not read -- a natural assumption, and
wrong. Every OpenStax file has both paragraphs, which is why reading .docx
directly surfaced a title that blocked the heading promotion.

Word declares most embedded images as application/octet-stream rather
than by their real type. Pandoc turns that into a ".so" extension and then
writes <embed> instead of <img>, giving a page that validates and shows
nothing. python-docx sets the content type from the file extension, so
producing the real-world case means rewriting [Content_Types].xml after
saving.

Copyright 2026 Robert Szarka

This program is free software: you can redistribute it and/or modify
it under the terms of the GNU General Public License as published by
the Free Software Foundation, either version 3 of the License, or
any later version.

This program is distributed in the hope that it will be useful,
but WITHOUT ANY WARRANTY; without even the implied warranty of
MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
GNU General Public License for more details.

You should have received a copy of the GNU General Public License
along with this program.  If not, see <https://www.gnu.org/licenses/>.
"""

import base64
import os
import re
import sys
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
FIXTURES = os.path.join(HERE, "fixtures")

try:
    import docx
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches
except ImportError:
    sys.exit("This needs python-docx:\n    pip3 install python-docx\n"
             "The built fixtures are committed, so the tests themselves "
             "do not.")

# A 1x1 PNG, enough to be an image without being worth storing.
PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAE"
    "hQGAhKmMIQAAAABJRU5ErkJggg==")


# --------------------------------------------------------------------------
# helpers
# --------------------------------------------------------------------------

def png_path():
    path = os.path.join(FIXTURES, "_dot.png")
    with open(path, "wb") as handle:
        handle.write(PNG)
    return path


def bold_cell(cell, text):
    cell.text = ""
    cell.paragraphs[0].add_run(text).bold = True


def mark_header_row(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))


def titled(d, text, style):
    """A paragraph in a named style, creating the style if the template
    lacks it. Pandoc resolves a pStyle to the style's name, so a style
    that only exists as an id is not recognized."""
    if style not in [s.name for s in d.styles]:
        d.styles.add_style(style, WD_STYLE_TYPE.PARAGRAPH)
    p = d.add_paragraph(text)
    p.style = d.styles[style]
    return p


def repack(path, edit):
    """Rewrite parts of a saved .docx. edit(name, bytes) -> bytes."""
    with zipfile.ZipFile(path) as archive:
        parts = {i.filename: archive.read(i.filename)
                 for i in archive.infolist()}
    parts = {name: edit(name, blob) for name, blob in parts.items()}
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        # [Content_Types].xml first, as the packaging convention requires.
        for name in sorted(parts, key=lambda n: n != "[Content_Types].xml"):
            out.writestr(name, parts[name])


def declare_media_as_octet_stream(name, blob):
    if name != "[Content_Types].xml":
        return blob
    text = blob.decode("utf-8")
    text = re.sub(r'(<Default Extension="png" ContentType=")[^"]*(")',
                  r"\1application/octet-stream\2", text)
    return text.encode("utf-8")


def set_first_image_description(description):
    def edit(name, blob):
        if name != "word/document.xml":
            return blob
        text = blob.decode("utf-8")
        text = re.sub(r'(<wp:docPr id="\d+" name="[^"]*")',
                      r'\1 descr="%s"' % description, text, count=1)
        return text.encode("utf-8")
    return edit


# --------------------------------------------------------------------------
# the fixtures
# --------------------------------------------------------------------------

def build_metadata(dot):
    """Title and Author paragraphs, and an H1 that is not the first block.

    Reproduces three things at once: the document metadata that blocked
    the heading promotion, the Author paragraph that produced a visible
    byline under every title, and a chapter-opening figure in a layout
    table, which is why finding the leading H1 means looking past the
    first block rather than at it.
    """
    d = docx.Document()
    titled(d, "Levels of Measurement", "Title")
    titled(d, "OpenStax", "Author")

    table = docx.Document  # noqa: F841  (readability only)
    t = d.add_table(rows=1, cols=1)
    t.rows[0].cells[0].paragraphs[0].add_run().add_picture(dot, width=Inches(2))
    d.add_paragraph("Figure 1.1 A diagram of something")

    d.add_heading("1.3 Levels of Measurement", level=1)
    d.add_paragraph("Some body text so the page is not empty.")
    path = os.path.join(FIXTURES, "metadata.docx")
    d.save(path)
    return path


def build_tables(dot):
    """A label Word stored as a one-item list, and a blank worksheet table.

    The label is the shape Markdown used to flatten, so reading .docx
    directly stopped pairing it with its table until the filter learned to
    unwrap it. The worksheet table is the shape the Markdown writer
    destroyed: every data cell empty, which came back as a paragraph break
    and reduced a ten-row table to one cell.
    """
    d = docx.Document()
    d.add_heading("Practice", level=1)

    label = d.add_paragraph("Table 1.1")
    label.style = d.styles["List Bullet"]
    t = d.add_table(rows=3, cols=2)
    for column, heading in enumerate(["Interval", "Frequency"]):
        bold_cell(t.rows[0].cells[column], heading)
    mark_header_row(t.rows[0])
    for row, values in enumerate([("0-9", "3"), ("10-19", "7")], start=1):
        for column, value in enumerate(values):
            t.rows[row].cells[column].text = value

    d.add_paragraph("Complete the table below.")
    d.add_paragraph("Table 1.2")
    blank = d.add_table(rows=6, cols=2)
    for column, heading in enumerate(["x", "P(x)"]):
        bold_cell(blank.rows[0].cells[column], heading)
    mark_header_row(blank.rows[0])
    # Every data cell empty: the trigger.

    path = os.path.join(FIXTURES, "tables.docx")
    d.save(path)
    return path


def build_media(dot, name, description):
    """An image whose description has a leading space, declared by Word as
    application/octet-stream.

    Two fixtures are built from this so a run covering both has two
    documents whose media is named identically inside the package. That is
    the collision that applied one document's alt text to another's image,
    because the filter sees the mediabag path before --extract-media adds
    the per-document directory.
    """
    d = docx.Document()
    d.add_heading(name.replace("-", " ").title(), level=1)
    d.add_paragraph().add_run().add_picture(dot, width=Inches(1))
    d.add_paragraph("Some text after the image.")
    # A second image with no description at all, so the alt-text report
    # has a row for this document. Without one there is nothing to check
    # the keys of, and the collision this guards against is invisible.
    d.add_paragraph().add_run().add_picture(dot, width=Inches(1))
    path = os.path.join(FIXTURES, name + ".docx")
    d.save(path)

    def edit(part, blob):
        blob = declare_media_as_octet_stream(part, blob)
        return set_first_image_description(description)(part, blob)
    repack(path, edit)
    return path


# An equation with a vertical bar in it. The Markdown intermediate lost
# three of these per page in the statistics book, because "|" is a table
# delimiter there and the writer had nowhere to put it. python-docx cannot
# create equations, so the OMML goes in as raw XML after saving -- which
# is also why this needs no copy of Word.
OMML = (
    '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/'
    '2006/math">'
    '<m:r><m:t>P(A|B) = </m:t></m:r>'
    '<m:f><m:num><m:r><m:t>P(A and B)</m:t></m:r></m:num>'
    '<m:den><m:r><m:t>P(B)</m:t></m:r></m:den></m:f>'
    '</m:oMath>')

# The detector looks for a run of four or more single-letter words, which
# is what spelling an identifier out produces. "cap a vertical bar cap b"
# does not qualify: its single letters are separated by words.
MATHSPEAK = "the value of s i g m a squared over n"


def build_math(dot):
    """An equation Word stored as OMML, and an image whose description is
    MathSpeak.

    MathSpeak is the notation maths speech engines use, and Word leaves it
    in the description when an equation was pasted as a picture. It spells
    identifiers out one letter at a time, which a screen reader then reads
    letter by letter. The filter rejoins runs of four or more single
    letters and says so, because the real fix is to re-author the equation
    in the source.
    """
    d = docx.Document()
    d.add_heading("Conditional probability", level=1)
    d.add_paragraph("The rule is ")
    d.add_paragraph("An equation stored as a picture follows.")
    d.add_paragraph().add_run().add_picture(dot, width=Inches(1))
    # A hair-thin image: Word's spacers, used for vertical rhythm and
    # carrying no meaning.
    d.add_paragraph().add_run().add_picture(dot, width=Inches(0.05))
    path = os.path.join(FIXTURES, "math.docx")
    d.save(path)

    def edit(part, blob):
        if part != "word/document.xml":
            return blob
        text = blob.decode("utf-8")
        text = text.replace("The rule is </w:t></w:r>",
                            "The rule is </w:t></w:r>" + OMML, 1)
        text = re.sub(r'(<wp:docPr id="\d+" name="[^"]*")',
                      r'\1 descr="%s"' % MATHSPEAK, text, count=1)
        return text.encode("utf-8")
    repack(path, edit)
    return path


def build_tables_b(dot):
    """A merged full-width row above the headers, and a table with no
    header signal at all.

    Word writes a table's title as a merged first row, which arrives as a
    header row spanning every column -- five tables in the statistics book
    look like this. The second table has no bold, no repeating row and no
    merged title, which is the shape nothing can classify and a person has
    to declare.
    """
    d = docx.Document()
    d.add_heading("More tables", level=1)

    t = d.add_table(rows=3, cols=3)
    merged = t.rows[0].cells[0].merge(t.rows[0].cells[2])
    merged.text = "Table 2.1 Sample results"
    for column, heading in enumerate(["Group", "Mean", "SD"]):
        bold_cell(t.rows[1].cells[column], heading)
    mark_header_row(t.rows[1])
    for column, value in enumerate(["A", "3.4", "0.8"]):
        t.rows[2].cells[column].text = value

    d.add_paragraph("And one with nothing to go on:")
    plain = d.add_table(rows=3, cols=2)
    for row, values in enumerate([("x", "y"), ("1", "2"), ("3", "4")]):
        for column, value in enumerate(values):
            plain.rows[row].cells[column].text = value

    path = os.path.join(FIXTURES, "tables-b.docx")
    d.save(path)
    return path


def main():
    os.makedirs(FIXTURES, exist_ok=True)
    dot = png_path()
    built = [
        build_metadata(dot),
        build_tables(dot),
        build_tables_b(dot),
        build_math(dot),
        build_media(dot, "media-a", " Alt text with a leading space"),
        build_media(dot, "media-b", "Alt text for the other document"),
    ]
    os.remove(dot)
    for path in built:
        print(f"  {os.path.basename(path):16} {os.path.getsize(path):6} bytes")
    return 0


if __name__ == "__main__":
    sys.exit(main())
