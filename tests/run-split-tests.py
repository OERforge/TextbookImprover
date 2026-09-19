#!/usr/bin/env python3
"""
run-split-tests.py -- check split-pages.py, and what reads its pieces.

    python3 tests/run-split-tests.py           # run every case
    python3 tests/run-split-tests.py --keep    # leave the output in place

Needs Pandoc 3.9 or later, as convert.sh does. The source is written as
Markdown here and turned into a .docx by Pandoc, because no fixture has
headings below its title and a chapter-shaped document is the whole
point. It then goes through the pre-pass and the filter the way
convert.sh sends it, and is cut from the filtered intermediate.

WHAT IS LOAD-BEARING HERE

A piece is the content between two cut headings, titled by the first,
with its own headings one level up and the cut heading's id kept as an
anchor. Names come from headings unless the sidecar says otherwise, and
a sidecar row that matches nothing is reported. Links between pieces are
rewritten to the piece the target moved to and links within a piece are
left alone. Every piece says where it came from, and both the packager
and the EPUB assembler read that: the pieces of a source group under it,
in reading order, without a contents tree saying so.

Copyright 2026 Robert Szarka

This program is free software: you can redistribute it and/or modify
it under the terms of the GNU General Public License as published by
the Free Software Foundation, either version 3 of the License, or
any later version.

This program is distributed in the hope that it will be useful,
but WITHOUT ANY WARRANTY; without even the implied warranty of
MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
GNU General Public License for more details.

You should have received a copy of the GNU General Public License
along with this program.  If not, see <https://www.gnu.org/licenses/>.
"""

import argparse
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
BIN = os.path.join(ROOT, "bin")
sys.path.insert(0, os.path.join(ROOT, "lib"))
import bookcontents  # noqa: E402


def load(path, name):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


cartridge = load(os.path.join(BIN, "build-cartridge.py"), "cartridge")

SOURCE = """\
---
title: Costs in the Long Run
---

# 7.5 Costs in the Long Run

The long run is when all costs vary. See [scale](#economies-of-scale).

## Choice of Production Technology

A firm can substitute machines for workers.

### Substituting machines for workers

Back to [the section](#choice-of-production-technology).

## Economies of Scale

Costs fall as output rises.

## Economies of Scale

A second section with the same heading, on purpose.

#

# Diseconomies

## Economies of Scale

The same heading under a different section.
"""


def run(arguments, cwd, environment=None, check=True):
    result = subprocess.run(arguments, cwd=cwd, env=environment,
                            capture_output=True, text=True,
                            stdin=subprocess.DEVNULL)
    if check and result.returncode:
        raise RuntimeError(f"{' '.join(arguments[:3])} failed:\n"
                           f"{result.stderr}")
    return result


def prepare(work, stem="chapter-7", source=SOURCE):
    """A filtered intermediate, as convert.sh leaves one."""
    os.makedirs(work, exist_ok=True)
    with open(os.path.join(work, "src.md"), "w", encoding="utf-8") as fh:
        fh.write(source)
    run(["pandoc", "src.md", "-o", stem + ".docx"], work)
    environment = dict(os.environ)
    environment.update({
        "TABLE_CAPTIONS": os.path.join(work, "table-captions.csv"),
        "IMAGE_ALT": os.path.join(work, "image-alt.csv"),
        "TABLE_HEADERS": os.path.join(work, "table-headers.csv"),
        "IMAGE_ALT_MISSING": os.path.join(work, "alt-missing.csv"),
        "TABLE_CAPTIONS_MISSING": os.path.join(work, "caps-missing.csv"),
        "SPACER_LOG": os.path.join(work, "spacers.csv"),
        "PROMOTE_H1_TO_TITLE": "always",
        "AUTHOR_BYLINE": "meta",
    })
    run(["pandoc", "-f", "docx", "-t", "json", stem + ".docx",
         "-o", stem + ".json",
         "--lua-filter", os.path.join(BIN, "media-extensions.lua"),
         "--extract-media", stem], work, environment)
    run(["pandoc", "-f", "json", "-t", "json", stem + ".json",
         "-o", stem + ".filtered.json",
         "--lua-filter", os.path.join(BIN, "figures-and-tables.lua")],
        work, environment)
    return os.path.join(work, stem + ".filtered.json")


def split(work, files, level=2, sidecar=None):
    arguments = ["python3", os.path.join(BIN, "split-pages.py"),
                 "--level", str(level),
                 "--sidecar", os.path.join(work, "page-names.csv"),
                 "--new", os.path.join(work, "page-names-new.csv"),
                 "--report", os.path.join(work, "page-names-report.csv")]
    if sidecar is not None:
        with open(os.path.join(work, "page-names.csv"), "w",
                  encoding="utf-8") as fh:
            fh.write(sidecar)
    result = run(arguments + files, work, check=False)
    return result


class Pieces:
    def __init__(self, work, result):
        self.work = work
        self.status = result.returncode
        self.stderr = result.stderr
        self.paths = result.stdout.split()
        self.stems = [os.path.basename(p)[:-len(".filtered.json")]
                      for p in self.paths]
        self.docs = {}
        for stem, path in zip(self.stems, self.paths):
            with open(path, encoding="utf-8") as fh:
                self.docs[stem] = json.load(fh)

    def text(self, stem):
        return json.dumps(self.docs[stem])

    def title(self, stem):
        meta = self.docs[stem]["meta"]["title"]
        return " ".join(("".join(i.get("c", " ") if i["t"] == "Str" else " "
                                 for i in meta["c"])).split())

    def headers(self, stem):
        return [b["c"][0] for b in self.docs[stem]["blocks"]
                if b["t"] == "Header"]

    def links(self, stem):
        return re.findall(r'"t": "Link", "c": \[\[.*?\], \[.*?\], \["([^"]*)"',
                          self.text(stem))

    def render(self, stem):
        """The page as convert.sh would write it."""
        head = os.path.join(self.work, "head.html")
        with open(head, "w", encoding="utf-8") as fh:
            fh.write("<style>/* test */</style>\n")
        environment = dict(os.environ, HEADER_INCLUDES_FILE=head)
        run(["pandoc", "-f", "json", "-t", "html5", stem + ".filtered.json",
             "-o", stem + ".html", "--standalone", "--ascii",
             "--lua-filter", os.path.join(BIN, "header-includes.lua"),
             "-M", "lang=en"], self.work, environment)
        with open(os.path.join(self.work, stem + ".html"),
                  encoding="utf-8") as fh:
            return fh.read()


# --------------------------------------------------------------------------
# the cases
# --------------------------------------------------------------------------

def case_pieces(work):
    """What a piece is."""
    source = prepare(work)
    out = Pieces(work, split(work, [source]))
    with open(os.path.join(work, "page-names-new.csv"),
              encoding="utf-8") as fh:
        new = fh.read()
    with open(os.path.join(work, "page-names-report.csv"),
              encoding="utf-8") as fh:
        report = fh.read()
    return [
        ("pieces in reading order, named by their headings",
         lambda: out.stems == [
             "chapter-7--7-5-costs-in-the-long-run",
             "chapter-7--choice-of-production-technology",
             "chapter-7--economies-of-scale",
             "chapter-7--economies-of-scale-2",
             "chapter-7--diseconomies--economies-of-scale"]),
        ("a repeated heading under another section takes its parent's name",
         lambda: "chapter-7--diseconomies--economies-of-scale" in out.stems),
        ("an empty heading cuts nothing",
         lambda: not any("part" in s for s in out.stems)),
        ("a heading with nothing of its own is a group, not a page",
         lambda: "chapter-7--diseconomies" not in out.stems
         and out.docs["chapter-7--diseconomies--economies-of-scale"]["meta"][
             "page-parents"]["c"] == [{"t": "MetaString", "c": "Diseconomies"}]),
        ("with two H1s the title stays a heading, so nothing precedes the "
         "first cut and there is no page named for the source",
         lambda: "chapter-7" not in out.stems
         and out.title("chapter-7--7-5-costs-in-the-long-run")
         == "7.5 Costs in the Long Run"
         and '"vary."' in out.text("chapter-7--7-5-costs-in-the-long-run")),
        ("a cut heading becomes its piece's title and leaves the body",
         lambda: out.title("chapter-7--economies-of-scale")
         == "Economies of Scale"
         and out.headers("chapter-7--economies-of-scale") == []),
        ("its id stays as an empty anchor at the top",
         lambda: out.docs["chapter-7--economies-of-scale"]["blocks"][0]
         == {"t": "Div", "c": [["economies-of-scale", ["page-title-anchor"],
                                []], []]}),
        ("a piece's own headings move up so it starts at h2",
         lambda: out.headers("chapter-7--choice-of-production-technology")
         == [2]),
        ("a link to a heading in another piece follows it there",
         lambda: "chapter-7--economies-of-scale.html#economies-of-scale"
         in out.links("chapter-7--7-5-costs-in-the-long-run")),
        ("a link within a piece is left alone",
         lambda: "#choice-of-production-technology"
         in out.links("chapter-7--choice-of-production-technology")),
        ("a repeated heading under the same parent gets a number and a "
         "warning",
         lambda: "repeats an earlier name" in out.stderr),
        ("each piece records its source, part, and position",
         lambda: out.docs["chapter-7--economies-of-scale"]["meta"][
             "page-part"]["c"] == "3/5"
         and out.docs["chapter-7--economies-of-scale"]["meta"][
             "page-position"]["c"] == "1.2"
         and out.docs["chapter-7--7-5-costs-in-the-long-run"]["meta"][
             "source-page"]["c"] == "chapter-7"),
        ("the prefilled rows give the whole name, and a position only "
         "where the heading path repeats",
         lambda: "chapter-7,7.5 Costs in the Long Run,Choice of Production "
         "Technology,,chapter-7--choice-of-production-technology" in new
         and "chapter-7,7.5 Costs in the Long Run,Economies of Scale,1.3,"
         "chapter-7--economies-of-scale-2" in new
         and "chapter-7,Diseconomies,Economies of Scale,,chapter-7--"
         "diseconomies--economies-of-scale" in new),
        ("the report lists every piece with its position and part",
         lambda: "chapter-7,,7.5 Costs in the Long Run,chapter-7--7-5-costs-"
         "in-the-long-run,1,1,5" in report
         and "economies-of-scale-2,1.3,4,5" in report),
    ]


def case_names(work):
    """The sidecar, and what it does with a row that matches nothing."""
    source = prepare(work)
    first = split(work, [source])
    prepare(work)                   # the split consumed the intermediate
    out = Pieces(work, split(work, [source], sidecar=(
        "source,parents,heading,position,name\n"
        "chapter-7,7.5 Costs in the Long Run,Choice of Production "
        "Technology,,7-1-technology\n"
        "chapter-7,7.5 Costs in the Long Run,Economies of Scale,1.3,"
        "7-3-scale-again\n"
        "chapter-7,Diseconomies,Economies of Scale,,7-4-scale\n"
        "chapter-7,,Gone Heading,,x\n")))
    bad = Pieces(os.path.join(work, "bad"), split(
        os.path.join(work, "bad"), [prepare(os.path.join(work, "bad"))],
        sidecar="source,parents,heading,position,name\n"
                "chapter-7,7.5 Costs in the Long Run,Economies of Scale,,"
                "no spaces allowed\n"))
    new_path = os.path.join(work, "page-names-new.csv")
    with open(new_path, encoding="utf-8") as fh:
        new = fh.read()
    return [
        ("a sidecar name is the whole page name",
         lambda: "7-1-technology" in out.stems),
        ("a row with a position names only the piece at that position",
         lambda: "7-3-scale-again" in out.stems
         and "chapter-7--economies-of-scale" in out.stems),
        ("a row is keyed on the headings above as well",
         lambda: "7-4-scale" in out.stems),
        ("a page the previous run wrote under another name is replaced",
         lambda: first.returncode == 0 and not os.path.exists(os.path.join(
             work, "chapter-7--choice-of-production-technology.filtered.json"))),
        ("a row for a heading the source does not have is reported",
         lambda: "no heading 'Gone Heading'" in out.stderr),
        ("only the unnamed pieces get prefilled rows",
         lambda: new.count("\n") == 3 and "Scale,1.2," in new),
        ("an unusable name is refused, said so, and the heading's used",
         lambda: bad.status == 0 and "not a usable name" in bad.stderr
         and "chapter-7--economies-of-scale" in bad.stems),
    ]


def case_levels(work):
    """Level 0 and a source with nothing to cut."""
    source = prepare(work)
    whole = Pieces(os.path.join(work, "whole"), split(work, [source], level=0))
    flat = prepare(os.path.join(work, "flat"), "notes",
                   "---\ntitle: Notes\n---\n\n# Notes\n\nJust a paragraph.\n")
    untouched = Pieces(os.path.join(work, "flat"),
                       split(os.path.join(work, "flat"), [flat], level=2))
    reserved = split(work, [source.replace("chapter-7", "a--b")], level=2)
    return [
        ("level 0 lists the source itself and writes nothing",
         lambda: whole.stems == ["chapter-7"]
         and not os.path.exists(os.path.join(work, "page-names-report.csv"))),
        ("a source with no heading at that level is left whole",
         lambda: untouched.stems == ["notes"]),
        ("a source whose name holds the separator is refused",
         lambda: reserved.returncode != 0 and "reserved" in reserved.stderr),
    ]


def case_readers(work):
    """What the packager and the EPUB assembler make of the pieces."""
    source = prepare(work)
    out = Pieces(work, split(work, [source]))
    titles, parts = {}, {}
    for stem in out.stems:
        page = out.render(stem)
        path = os.path.join(work, stem + ".html")
        titles[stem] = cartridge.page_title(path, stem)
        origin = cartridge.page_provenance(path)
        if origin:
            parts[stem] = origin
    guessed = bookcontents.guess_contents(out.stems, None, titles, parts)
    by_name = bookcontents.guess_contents(
        sorted(out.stems), None, titles)     # no provenance: by name

    with open(os.path.join(work, "project.yaml"), "w", encoding="utf-8") as fh:
        fh.write("project:\n  identifier: org.example.costs\n  title: Costs\n")
    with open(os.path.join(work, "conversion.yaml"), "w",
              encoding="utf-8") as fh:
        fh.write("targets:\n  epub:\n    format: epub3\n")
    built = run(["python3", os.path.join(BIN, "build-epub.py"), "-d", work],
                work, check=False)
    epub_path = os.path.join(work, "epub", "org.example.costs.epub")
    chapters, nav = {}, ""
    if os.path.exists(epub_path):
        with zipfile.ZipFile(epub_path) as archive:
            for name in archive.namelist():
                if "/text/ch" in name:
                    chapters[name] = archive.read(name).decode("utf-8")
                elif name.endswith("nav.xhtml"):
                    nav = archive.read(name).decode("utf-8")
    body = "\n".join(chapters.values())
    return [
        ("the page head carries the provenance, and the packager reads it",
         lambda: parts["chapter-7--diseconomies--economies-of-scale"]
         == ("chapter-7", 5, ["Diseconomies"], "2.1")),
        ("the packager groups the pieces under the source and their "
         "headings, a heading's own page opening its group",
         lambda: guessed == [
             {"title": "7.5 Costs in the Long Run", "items": out.stems[:4]},
             {"title": "Diseconomies", "items": [out.stems[4]]}]),
        ("without provenance the pieces still group, by name",
         lambda: by_name and all(isinstance(n, str) for n in by_name)
         and set(by_name) == set(out.stems)),
        ("the EPUB builds from the pieces, and not from the source too",
         lambda: built.returncode == 0 and len(chapters) == 7
         and not os.path.exists(os.path.join(work, "chapter-7.filtered.json"))),
        ("and its nav shows the source over its pieces",
         lambda: nav.index("7.5 Costs in the Long Run")
         < nav.index("Choice of Production Technology")),
        ("a link between pieces resolves inside the book",
         lambda: re.search(r'href="ch00\d\.xhtml#page-chapter-7--economies-'
                           r'of-scale--economies-of-scale"', body)),
    ]


def case_restyle(work):
    """util/restyle-headings.py: the source repair a Title-styled book
    needs before any of the above can see its top level."""
    import docx
    os.makedirs(work, exist_ok=True)
    d = docx.Document()
    d.add_paragraph("", style="Title")               # Word's cruft
    d.add_paragraph("Module 1: Basics", style="Title")
    d.add_paragraph("Intro text")
    d.add_paragraph("What is Java?", style="Heading 1")
    d.add_paragraph("Java Goals", style="Heading 2")
    d.add_paragraph("", style="Heading 2")             # empty, and remapped
    d.add_paragraph("Module 2: More", style="Title")
    d.add_paragraph("Body text", style="Body Text")
    path = os.path.join(work, "book.docx")
    d.save(path)
    tool = os.path.join(ROOT, "util", "restyle-headings.py")
    report = run(["python3", tool, path], work)
    no_toc = run(["python3", tool, path, "--from-toc"], work, check=False)
    out = os.path.join(work, "restyled.docx")
    done = run(["python3", tool, path, "--promote", "Title", "-o", out], work)
    kept = os.path.join(work, "kept.docx")
    run(["python3", tool, path, "--promote", "Title", "--keep-empty",
         "-o", kept], work)
    run(["pandoc", "-f", "docx", "-t", "json", out, "-o", "restyled.json"],
        work)
    with open(os.path.join(work, "restyled.json"), encoding="utf-8") as fh:
        doc = json.load(fh)
    headers = [(b["c"][0], "".join(i.get("c", " ") if i["t"] == "Str"
                                   else " " for i in b["c"][2]))
               for b in doc["blocks"] if b["t"] == "Header"]
    with zipfile.ZipFile(kept) as archive:
        kept_xml = archive.read("word/document.xml").decode("utf-8")
    with zipfile.ZipFile(path) as archive:
        original = archive.namelist()
    with zipfile.ZipFile(out) as archive:
        restyled = archive.namelist()

    # A book with several H1 sections and no heading over them.
    flat = docx.Document()
    flat.add_paragraph("What is Java?", style="Heading 1")
    flat.add_paragraph("Java Goals", style="Heading 2")
    flat.add_paragraph("Comments", style="Heading 1")
    flat_path = os.path.join(work, "flat.docx")
    flat.save(flat_path)
    titled = os.path.join(work, "titled.docx")
    run(["python3", tool, flat_path, "--demote", "--title", "Intro & Java",
         "-o", titled], work)
    run(["pandoc", "-f", "docx", "-t", "json", titled, "-o", "titled.json"],
        work)
    with open(os.path.join(work, "titled.json"), encoding="utf-8") as fh:
        tdoc = json.load(fh)
    theaders = [(b["c"][0], "".join(i.get("c", " ") if i["t"] == "Str"
                                    else " " for i in b["c"][2]))
                for b in tdoc["blocks"] if b["t"] == "Header"]
    with zipfile.ZipFile(titled) as archive:
        core = archive.read("docProps/core.xml").decode("utf-8")
    return [
        ("--demote --title puts one H1 over demoted sections",
         lambda: theaders == [(1, "Intro & Java"), (2, "What is Java?"),
                              (3, "Java Goals"), (2, "Comments")]),
        ("and records it as the document's dc:title, escaped",
         lambda: "<dc:title>Intro &amp; Java</dc:title>" in core),
        ("with no option the tool reports and writes nothing",
         lambda: "Title" in report.stdout and "Nothing changed" in report.stdout
         and not os.path.exists(os.path.join(work, "book.docx.tmp"))),
        ("--from-toc refuses a document with no TOC field",
         lambda: no_toc.returncode != 0 and "no TOC field" in no_toc.stderr),
        ("--promote moves every heading down one, all at once",
         lambda: headers == [(1, "Module 1: Basics"), (2, "What is Java?"),
                             (3, "Java Goals"), (1, "Module 2: More")]),
        ("an empty paragraph of a remapped style is dropped, and said so",
         lambda: "empty Title paragraph(s) dropped" in done.stdout
         and "empty Heading2 paragraph(s) dropped" in done.stdout),
        ("--keep-empty keeps it",
         lambda: kept_xml.count('w:val="Heading1"') == 3),
        ("the metadata title is gone with the Title style, as it should be",
         lambda: "title" not in doc["meta"]),
        ("every other part is carried over",
         lambda: original == restyled),
    ]


CASES = [
    ("restyling a Title-styled source", case_restyle),
    ("what a piece is", case_pieces),
    ("the page-names sidecar", case_names),
    ("levels and refusals", case_levels),
    ("the packager and the assembler read the pieces", case_readers),
]


# --------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="Check split-pages.py and what reads its pieces.")
    parser.add_argument("--keep", action="store_true",
                        help="leave the output in place")
    arguments = parser.parse_args()

    if shutil.which("pandoc") is None:
        sys.exit("pandoc is not on the path.")
    version = subprocess.run(["pandoc", "--version"], capture_output=True,
                             text=True).stdout.split()[1]
    if tuple(int(p) for p in re.findall(r"\d+", version)[:3]) < (3, 9):
        sys.exit(f"Pandoc {version} is too old; these tests need 3.9 or "
                 "later, as convert.sh does.")

    work = tempfile.mkdtemp(prefix="split-tests-")
    failed = 0
    try:
        for label, case in CASES:
            directory = os.path.join(work, re.sub(r"[^\w-]+", "-", label))
            try:
                checks = case(directory)
            except Exception as exc:
                print(f"  ERROR {label}: {exc}")
                failed += 1
                continue
            for name, predicate in checks:
                try:
                    passed = predicate()
                except Exception as exc:
                    passed, name = False, f"{name}  ({exc})"
                print(("  ok    " if passed else "  FAIL  ") + name)
                failed += not passed
    finally:
        if arguments.keep:
            print(f"\nOutput left in {work}")
        else:
            shutil.rmtree(work, ignore_errors=True)

    print(f"\n{failed} check(s) failed across {len(CASES)} case(s)"
          if failed else "\nall split checks passed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
